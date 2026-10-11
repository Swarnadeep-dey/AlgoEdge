"""Formatted Excel reports for AlgoEdge (backtest, optimizer, screener).

The three scripts (Backtester.py, Optimeser.py, Screeners.py) hand their results to
the functions below, so all workbook styling lives in one place.

Compatibility note: Dashboard.py reads some sheets by name and column header
('Executive Summary', 'Equity Curve', 'Benchmark', 'Trade Log', and the FIRST sheet of
Optimization.xlsx). Those sheets keep their header in row 1 with the same column names,
so changing the look here can never break the dashboard.
"""
import math

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import AreaChart, BarChart, LineChart, Reference
from openpyxl.chart.axis import DateAxis
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

# ----------------------------------------------------------------------------- style
FONT = "Arial"
NAVY, BLUE, ORANGE, RED, GREEN = "1F4E78", "2A78D6", "EB6834", "E34948", "0CA30C"
INK, MUTED, CARD, LINE, GRID = "1A1A19", "6B6A66", "F3F5F8", "D4D8DE", "E1E0D9"
RUPEE = "[$₹-en-IN]#,##0.00"
PCT_TRADE = '+0.00"%";-0.00"%";0.00"%"'

THIN = Side(style="thin", color=LINE)
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def F(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, italic=italic, color=color)


def fill(hex_):
    return PatternFill("solid", start_color=hex_, end_color=hex_)


def py(v):
    """numpy / pandas scalars -> plain Python so openpyxl writes the right cell type."""
    if v is None:
        return None
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime()
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        try:
            v = v.item()
        except Exception:
            pass
    if isinstance(v, float) and math.isnan(v):
        return None
    return v


def header_row(ws, row, c1, c2, height=22):
    for c in range(c1, c2 + 1):
        cell = ws.cell(row, c)
        cell.font = F(10, True, "FFFFFF")
        cell.fill = fill(NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX
    ws.row_dimensions[row].height = height


def title_block(ws, col, title, subtitle):
    ws.cell(1, col, title).font = F(22, True, NAVY)
    ws.row_dimensions[1].height = 34
    ws.cell(2, col, subtitle).font = F(10, False, MUTED)
    ws.row_dimensions[2].height = 18


def section_label(ws, row, col, text):
    c = ws.cell(row, col, text)
    c.font = F(12, True, NAVY)
    ws.row_dimensions[row].height = 22


def card(ws, top, left, span, label, value, note=None, number_format=None):
    """Three-row KPI card: small label, large value, small note. `value`/`note` may be formulas."""
    last = left + span - 1
    heights = (17, 32, 17)
    for i in range(3):
        r = top + i
        a = ws.cell(r, left)
        a.fill = fill(CARD)
        a.alignment = Alignment(horizontal="center", vertical="center")
        a.border = Border(left=THIN, right=THIN,
                          top=THIN if i == 0 else None, bottom=THIN if i == 2 else None)
        ws.row_dimensions[r].height = heights[i]
        ws.merge_cells(start_row=r, start_column=left, end_row=r, end_column=last)
    lab = ws.cell(top, left, label)
    lab.font = F(9, False, MUTED)
    val = ws.cell(top + 1, left, value)
    val.font = F(18, True, INK)
    if number_format:
        val.number_format = number_format
    nt = ws.cell(top + 2, left, note)
    nt.font = F(9, False, MUTED)


def style_chart(ch, title, width, height, legend="b"):
    ch.title = title
    ch.width, ch.height = width, height
    if legend is None:
        ch.legend = None
    else:
        ch.legend.position = legend
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.y_axis.majorGridlines.spPr = GraphicalProperties(ln=LineProperties(solidFill=GRID))


def color_series(s, color, width=None, fill_it=True):
    if fill_it:
        s.graphicalProperties.solidFill = color
    s.graphicalProperties.line.solidFill = color
    if width:
        s.graphicalProperties.line.width = width


def date_axis(ch, fmt="mmm-yy", unit=3):
    ch.x_axis = DateAxis(crossAx=100)
    ch.x_axis.number_format = fmt
    ch.x_axis.majorTimeUnit = "months"
    ch.x_axis.majorUnit = unit
    ch.x_axis.tickLblPos = "low"
    ch.x_axis.delete = False
    ch.y_axis.crossAx = 500
    ch.y_axis.delete = False


def landscape_fit(ws, one_page=False):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def nice_floor(x, step):
    return math.floor(x / step) * step


def nice_ceil(x, step):
    return math.ceil(x / step) * step


# ----------------------------------------------------------------------------- backtest
METRIC_HELP = {
    "Initial Capital": "Starting cash for the simulation.",
    "Final Portfolio Value": "Cash plus open positions at the end of the test window.",
    "CAGR (%)": "Compound annual growth rate of the portfolio.",
    "Annualized Volatility (%)": "Standard deviation of daily returns, scaled to a year.",
    "Sharpe Ratio (Rf=5%)": "Excess return over the 5% risk-free rate, per unit of volatility.",
    "Sortino Ratio": "Like Sharpe, but only penalises downside volatility.",
    "Max Drawdown (%)": "Largest peak-to-trough fall in portfolio value.",
    "Profit Factor": "Gross profit divided by gross loss across all trades.",
    "Strategy Beta": "Sensitivity to NIFTYBEES daily moves (1.0 = moves with the market).",
    "Strategy Alpha (Annualized)": "CAPM alpha: return not explained by market exposure.",
}


def write_backtest_report(path, df_equity, df_trades, df_metrics, df_bench, meta):
    """meta keys: rules (str), generated (str), test_window (str), benchmark (str)."""
    wb = Workbook()
    ws_d = wb.active
    ws_d.title = "Dashboard"
    ws_a = wb.create_sheet("Trade Analytics")
    ws_s = wb.create_sheet("Executive Summary")
    ws_e = wb.create_sheet("Equity Curve")
    has_bench = df_bench is not None and not df_bench.empty
    ws_b = wb.create_sheet("Benchmark") if has_bench else None
    ws_t = wb.create_sheet("Trade Log")
    for ws, color in ((ws_d, NAVY), (ws_a, NAVY), (ws_s, BLUE), (ws_e, BLUE), (ws_b, BLUE), (ws_t, BLUE)):
        if ws is not None:
            ws.sheet_properties.tabColor = color

    n_eq = len(df_equity)
    n_tr = len(df_trades)

    # ---- Executive Summary (header in row 1: Dashboard.py reads Metric / Value) ----
    ws_s.append(["Metric", "Value", "What it means"])
    header_row(ws_s, 1, 1, 3)
    metric_row = {}
    for i, (m, v) in enumerate(zip(df_metrics["Metric"], df_metrics["Value"]), start=2):
        ws_s.append([m, v, METRIC_HELP.get(m, "")])
        metric_row[m] = i
        ws_s.cell(i, 1).font = F(10, True)
        ws_s.cell(i, 2).font = F(10)
        ws_s.cell(i, 2).alignment = Alignment(horizontal="center")
        ws_s.cell(i, 3).font = F(10, False, MUTED)
        for c in (1, 2, 3):
            ws_s.cell(i, c).border = BOX
            if i % 2 == 1:
                ws_s.cell(i, c).fill = fill(CARD)
    ws_s.column_dimensions["A"].width = 32
    ws_s.column_dimensions["B"].width = 20
    ws_s.column_dimensions["C"].width = 70
    ws_s.freeze_panes = "A2"

    # ---- Equity Curve (Date, Portfolio Equity, Drawdown) ----
    ws_e.append(["Date", "Portfolio Equity", "Drawdown"])
    header_row(ws_e, 1, 1, 3)
    for i, (d, v) in enumerate(zip(df_equity["Date"], df_equity["Portfolio Equity"]), start=2):
        ws_e.cell(i, 1, py(d)).number_format = "yyyy-mm-dd"
        ws_e.cell(i, 1).alignment = Alignment(horizontal="center")
        ws_e.cell(i, 2, py(v)).number_format = RUPEE
        ws_e.cell(i, 3, f"=B{i}/MAX(B$2:B{i})-1").number_format = "0.00%"
        for c in (1, 2, 3):
            ws_e.cell(i, c).font = F(10)
    ws_e.column_dimensions["A"].width = 14
    ws_e.column_dimensions["B"].width = 22
    ws_e.column_dimensions["C"].width = 14
    ws_e.freeze_panes = "A2"

    # ---- Benchmark (Date, Strategy Equity, Benchmark Equity) ----
    if has_bench:
        df_cmp = df_equity.assign(Date=pd.to_datetime(df_equity["Date"])).merge(
            df_bench.assign(Date=pd.to_datetime(df_bench["Date"])), on="Date", how="left").ffill().bfill()
        df_cmp.columns = ["Date", "Strategy Equity", "Benchmark Equity"]
        ws_b.append(list(df_cmp.columns))
        header_row(ws_b, 1, 1, 3)
        for i, row in enumerate(df_cmp.itertuples(index=False), start=2):
            ws_b.cell(i, 1, py(row[0])).number_format = "yyyy-mm-dd"
            ws_b.cell(i, 1).alignment = Alignment(horizontal="center")
            ws_b.cell(i, 2, py(row[1])).number_format = RUPEE
            ws_b.cell(i, 3, py(row[2])).number_format = RUPEE
            for c in (1, 2, 3):
                ws_b.cell(i, c).font = F(10)
        for col, w in (("A", 14), ("B", 20), ("C", 20)):
            ws_b.column_dimensions[col].width = w
        ws_b.freeze_panes = "A2"

    # ---- Trade Log ----
    tr_cols = ["Ticker", "Entry Date", "Exit Date", "Days Held", "Return (%)", "Profit (INR)", "Win", "Exit Reason"]
    ws_t.append(tr_cols)
    header_row(ws_t, 1, 1, len(tr_cols))
    if n_tr:
        days = (pd.to_datetime(df_trades["Exit Date"]) - pd.to_datetime(df_trades["Entry Date"])).dt.days
        for i, (row, dh) in enumerate(zip(df_trades.itertuples(index=False), days), start=2):
            vals = [row[df_trades.columns.get_loc("Ticker")], row[df_trades.columns.get_loc("Entry Date")],
                    row[df_trades.columns.get_loc("Exit Date")], dh, row[df_trades.columns.get_loc("Return (%)")],
                    row[df_trades.columns.get_loc("Profit (INR)")], row[df_trades.columns.get_loc("Win")],
                    row[df_trades.columns.get_loc("Exit Reason")]]
            for c, v in enumerate(vals, start=1):
                cell = ws_t.cell(i, c, py(v))
                cell.font = F(10)
                cell.border = BOX
                if c in (2, 3):
                    cell.number_format = "yyyy-mm-dd"
                    cell.alignment = Alignment(horizontal="center")
                elif c == 4:
                    cell.number_format = "0"
                    cell.alignment = Alignment(horizontal="center")
                elif c == 5:
                    cell.number_format = PCT_TRADE
                elif c == 6:
                    cell.number_format = RUPEE
                elif c in (7, 8):
                    cell.alignment = Alignment(horizontal="center")
                if c == 7:
                    win = bool(py(v))
                    cell.fill = fill("C6EFCE" if win else "FFC7CE")
                    cell.font = F(10, False, "006100" if win else "9C0006")
        last = n_tr + 1
        tot = last + 1
        ws_t.cell(tot, 1, "TOTAL / AVG:").font = F(10, True)
        s = ws_t.cell(tot, 6, f"=SUM(F2:F{last})")
        s.number_format, s.font = RUPEE, F(10, True)
        a = ws_t.cell(tot, 5, f"=AVERAGE(E2:E{last})")
        a.number_format, a.font = PCT_TRADE, F(10, True)
        ws_t.auto_filter.ref = f"A1:H{last}"
    for col, w in (("A", 16), ("B", 14), ("C", 14), ("D", 11), ("E", 13), ("F", 18), ("G", 8), ("H", 13)):
        ws_t.column_dimensions[col].width = w
    ws_t.freeze_panes = "A2"

    # ---- Dashboard sheet ----
    ws_d.sheet_view.showGridLines = False
    ws_d.column_dimensions["A"].width = 2
    card_lefts = [2, 6, 10, 14, 18]
    for c in range(2, 21):
        ws_d.column_dimensions[L(c)].width = 1.5 if (c - 1) % 4 == 0 else 11
    title_block(ws_d, 2, "AlgoEdge Backtest Report",
                f"Test window {meta['test_window']}  |  Generated {meta['generated']}")
    chip = ws_d.cell(3, 2, meta["rules"] + (f"  |  Benchmark: {meta['benchmark']}" if has_bench
                                            else "  |  Benchmark data unavailable this run"))
    chip.font = F(9, False, MUTED)
    chip.fill = fill(CARD)
    chip.alignment = Alignment(horizontal="left", vertical="center", indent=1)
    ws_d.merge_cells(start_row=3, start_column=2, end_row=3, end_column=20)
    ws_d.row_dimensions[3].height = 20
    ws_d.row_dimensions[4].height = 8

    def mref(name):
        return f"='Executive Summary'!B{metric_row[name]}" if name in metric_row else "n/a"

    last_e = n_eq + 1
    trades_win = (f"=TEXT(COUNTIF('Trade Log'!E2:E{n_tr + 1},\">0\")/COUNT('Trade Log'!E2:E{n_tr + 1}),\"0%\")&\" win rate\""
                  if n_tr else "no completed trades")
    row1 = [
        ("Final value", mref("Final Portfolio Value"),
         f"=\"Return: \"&TEXT('Equity Curve'!B{last_e}/'Equity Curve'!B2-1,\"+0.00%;-0.00%\")"),
        ("CAGR", mref("CAGR (%)"), "annual growth rate"),
        ("Max drawdown", mref("Max Drawdown (%)"), "worst peak-to-trough fall"),
        ("Sharpe ratio", mref("Sharpe Ratio (Rf=5%)"), "vs 5% risk-free rate"),
        ("Profit factor", mref("Profit Factor"), "gross profit / gross loss"),
    ]
    row2 = [
        ("Volatility (ann.)", mref("Annualized Volatility (%)"), "std dev of daily returns"),
        ("Sortino ratio", mref("Sortino Ratio"), "penalises downside only"),
        ("Beta", mref("Strategy Beta"), f"vs {meta['benchmark']}"),
        ("Alpha (annualized)", mref("Strategy Alpha (Annualized)"), "return beyond market exposure"),
        ("Trades", f"=COUNT('Trade Log'!E2:E{n_tr + 1})" if n_tr else 0, trades_win),
    ]
    for (label, val, note), left in zip(row1, card_lefts):
        card(ws_d, 5, left, 3, label, val, note)
    ws_d.row_dimensions[8].height = 8
    for (label, val, note), left in zip(row2, card_lefts):
        card(ws_d, 9, left, 3, label, val, note, number_format="0" if label == "Trades" else None)
    ws_d.row_dimensions[12].height = 10

    # chart 1: strategy vs benchmark
    section_label(ws_d, 13, 2, "Portfolio value vs benchmark")
    ch = LineChart()
    date_axis(ch)
    if has_bench:
        ch.add_data(Reference(ws_b, min_col=2, max_col=3, min_row=1, max_row=len(df_cmp) + 1), titles_from_data=True)
        ch.set_categories(Reference(ws_b, min_col=1, min_row=2, max_row=len(df_cmp) + 1))
        lo = min(df_cmp["Strategy Equity"].min(), df_cmp["Benchmark Equity"].min())
        hi = max(df_cmp["Strategy Equity"].max(), df_cmp["Benchmark Equity"].max())
        colors = (BLUE, ORANGE)
    else:
        ch.add_data(Reference(ws_e, min_col=2, min_row=1, max_row=last_e), titles_from_data=True)
        ch.set_categories(Reference(ws_e, min_col=1, min_row=2, max_row=last_e))
        lo, hi = df_equity["Portfolio Equity"].min(), df_equity["Portfolio Equity"].max()
        colors = (BLUE,)
    style_chart(ch, "Portfolio value (₹)" + (" vs NIFTYBEES buy & hold" if has_bench else ""), 33, 9,
                legend="b" if has_bench else None)
    ch.y_axis.number_format = "#,##0"
    pad = max((hi - lo) * 0.08, 1000)
    ch.y_axis.scaling.min = nice_floor(lo - pad, 5000)
    ch.y_axis.scaling.max = nice_ceil(hi + pad, 5000)
    for s, c in zip(ch.series, colors):
        color_series(s, c, width=25000, fill_it=False)
        s.smooth = False
    ws_d.add_chart(ch, "B14")
    if has_bench:
        lb = len(df_cmp) + 1
        cap = ws_d.cell(32, 2,
                        f"=\"Strategy \"&TEXT(Benchmark!B{lb}/Benchmark!B2-1,\"+0.00%;-0.00%\")&\" vs {meta['benchmark']} \"&"
                        f"TEXT(Benchmark!C{lb}/Benchmark!C2-1,\"+0.00%;-0.00%\")&\" over the same window. "
                        f"The strategy is invested only part of the time, so it lags in strong rallies.\"")
        cap.font = F(9, False, MUTED, italic=True)

    # chart 2: drawdown
    section_label(ws_d, 34, 2, "Drawdown")
    dd = AreaChart()
    date_axis(dd)
    dd.add_data(Reference(ws_e, min_col=3, min_row=1, max_row=last_e), titles_from_data=True)
    dd.set_categories(Reference(ws_e, min_col=1, min_row=2, max_row=last_e))
    style_chart(dd, "Drawdown from previous peak", 33, 6.5, legend=None)
    dd.y_axis.number_format = "0%"
    dd_min = (df_equity["Portfolio Equity"] / df_equity["Portfolio Equity"].cummax() - 1).min()
    dd.y_axis.scaling.max = 0
    dd.y_axis.scaling.min = min(nice_floor(dd_min, 0.01) - 0.005, -0.01)
    dd.series[0].graphicalProperties.solidFill = "F2B8B8"
    dd.series[0].graphicalProperties.line.solidFill = RED
    ws_d.add_chart(dd, "B35")
    landscape_fit(ws_d, one_page=True)

    # ---- Trade Analytics ----
    ws_a.sheet_view.showGridLines = False
    ws_a.column_dimensions["A"].width = 2
    for col, w in (("B", 24), ("C", 14), ("D", 14), ("E", 14), ("F", 3)):
        ws_a.column_dimensions[col].width = w
    title_block(ws_a, 2, "Trade analytics", "All figures are formulas on the Trade Log sheet, so they update with it.")
    if not n_tr:
        ws_a.cell(4, 2, "No completed trades in this window.").font = F(11, False, MUTED)
        landscape_fit(ws_a)
        wb.active = 0
        wb.save(path)
        return

    last_t = n_tr + 1
    rng = lambda col: f"'Trade Log'!${col}$2:${col}${last_t}"
    TICK, DAYS, RET, PNL, REASON = rng("A"), rng("D"), rng("E"), rng("F"), rng("H")

    # how trades ended
    section_label(ws_a, 4, 2, "How trades ended")
    for c, h in enumerate(["Exit", "Trades", "Avg return (%)"], start=2):
        ws_a.cell(5, c, h)
    header_row(ws_a, 5, 2, 4)
    reasons = [("TP", "Take-profit"), ("SL", "Stop-loss"), ("RSI", "RSI exit")]
    for i, (code, name) in enumerate(reasons, start=6):
        ws_a.cell(i, 2, name).font = F(10, True)
        ws_a.cell(i, 3, f'=COUNTIF({REASON},"{code}")').font = F(10)
        ws_a.cell(i, 4, f'=IFERROR(AVERAGEIF({REASON},"{code}",{RET}),0)').font = F(10)
        ws_a.cell(i, 4).number_format = PCT_TRADE
        for c in (2, 3, 4):
            ws_a.cell(i, c).border = BOX
        ws_a.cell(i, 3).alignment = Alignment(horizontal="center")
        ws_a.cell(i, 4).alignment = Alignment(horizontal="center")
    ex = BarChart()
    ex.type = "bar"
    ex.add_data(Reference(ws_a, min_col=3, min_row=5, max_row=8), titles_from_data=True)
    ex.set_categories(Reference(ws_a, min_col=2, min_row=6, max_row=8))
    style_chart(ex, "How trades ended", 15, 6, legend=None)
    ex.x_axis.scaling.orientation = "maxMin"
    ex.y_axis.crosses = "max"
    ex.gapWidth = 60
    color_series(ex.series[0], BLUE)
    ex.dataLabels = DataLabelList()
    ex.dataLabels.showVal = True
    ex.dataLabels.showSerName = ex.dataLabels.showCatName = ex.dataLabels.showLegendKey = False
    ws_a.add_chart(ex, "G4")

    # summary
    section_label(ws_a, 11, 2, "Trade summary")
    summary = [
        ("Total trades", f"=COUNT({RET})", "0"),
        ("Win rate", f"=COUNTIF({RET},\">0\")/COUNT({RET})", "0.0%"),
        ("Average win", f"=IFERROR(AVERAGEIF({RET},\">0\"),0)", PCT_TRADE),
        ("Average loss", f"=IFERROR(AVERAGEIF({RET},\"<=0\"),0)", PCT_TRADE),
        ("Best trade", f"=MAX({RET})", PCT_TRADE),
        ("Worst trade", f"=MIN({RET})", PCT_TRADE),
        ("Average per trade", f"=AVERAGE({RET})", PCT_TRADE),
        ("Average holding period", f"=AVERAGE({DAYS})", '0" days"'),
        ("Total profit", f"=SUM({PNL})", RUPEE),
    ]
    for i, (lab, f, nf) in enumerate(summary, start=12):
        a = ws_a.cell(i, 2, lab)
        a.font = F(10, True)
        b = ws_a.cell(i, 3, f)
        b.number_format = nf
        b.font = F(10)
        b.alignment = Alignment(horizontal="center")
        ws_a.merge_cells(start_row=i, start_column=3, end_row=i, end_column=4)
        for c in (2, 3, 4):
            ws_a.cell(i, c).border = BOX
        if i % 2 == 1:
            a.fill = fill(CARD)
            b.fill = fill(CARD)

    # return distribution
    rets = pd.to_numeric(df_trades["Return (%)"])
    width = 1 if (math.ceil(rets.max()) - math.floor(rets.min())) <= 30 else 2
    lo_b = int(nice_floor(math.floor(rets.min()), width))
    hi_b = int(math.floor(rets.max())) + width
    bins = list(range(lo_b, hi_b, width))
    s0 = 23
    section_label(ws_a, s0, 2, "Distribution of trade returns")
    ws_a.cell(s0 + 1, 2, "Bin width (%)").font = F(10, True)
    wcell = ws_a.cell(s0 + 1, 3, width)
    wcell.font = F(10, False, "0000FF")  # input cell: blue
    wcell.alignment = Alignment(horizontal="center")
    ws_a.cell(s0 + 1, 4, "(edit to re-bin)").font = F(9, False, MUTED, italic=True)
    hdr = s0 + 2
    for c, h in enumerate(["Bin start (%)", "Losses", "Wins"], start=2):
        ws_a.cell(hdr, c, h)
    header_row(ws_a, hdr, 2, 4)
    for k, b in enumerate(bins):
        r = hdr + 1 + k
        ws_a.cell(r, 2, b).number_format = '0"%"'
        cnt = f'COUNTIFS({RET},">="&$B{r},{RET},"<"&($B{r}+$C${s0 + 1}))'
        ws_a.cell(r, 3, f"=IF($B{r}<0,{cnt},0)")
        ws_a.cell(r, 4, f"=IF($B{r}>=0,{cnt},0)")
        for c in (2, 3, 4):
            ws_a.cell(r, c).font = F(10)
            ws_a.cell(r, c).alignment = Alignment(horizontal="center")
            ws_a.cell(r, c).border = BOX
    h_end = hdr + len(bins)
    hist = BarChart()
    hist.type = "col"
    hist.grouping = "stacked"
    hist.overlap = 100
    hist.gapWidth = 15
    hist.add_data(Reference(ws_a, min_col=3, max_col=4, min_row=hdr, max_row=h_end), titles_from_data=True)
    hist.set_categories(Reference(ws_a, min_col=2, min_row=hdr + 1, max_row=h_end))
    style_chart(hist, "Return per trade, net of costs", 20, 8, legend="b")
    color_series(hist.series[0], RED)
    color_series(hist.series[1], GREEN)
    ws_a.add_chart(hist, f"G{s0}")

    # profit by stock
    by_stock = df_trades.groupby("Ticker")["Profit (INR)"].sum().sort_values(ascending=False)
    p0 = max(h_end + 3, s0 + 20)
    section_label(ws_a, p0, 2, "Profit by stock")
    for c, h in enumerate(["Ticker", "Profit (₹)", "Gain", "Loss"], start=2):
        ws_a.cell(p0 + 1, c, h)
    ws_a.column_dimensions["E"].width = 14
    header_row(ws_a, p0 + 1, 2, 5)
    for k, (tk, _) in enumerate(by_stock.items()):
        r = p0 + 2 + k
        ws_a.cell(r, 2, tk).font = F(10, True)
        ws_a.cell(r, 3, f"=SUMIF({TICK},B{r},{PNL})").number_format = RUPEE
        ws_a.cell(r, 4, f"=IF(C{r}>=0,C{r},0)").number_format = RUPEE
        ws_a.cell(r, 5, f"=IF(C{r}<0,C{r},0)").number_format = RUPEE
        for c in (2, 3, 4, 5):
            ws_a.cell(r, c).border = BOX
            if c > 2:
                ws_a.cell(r, c).font = F(10)
    p_end = p0 + 1 + len(by_stock)
    pb = BarChart()
    pb.type = "bar"
    pb.grouping = "stacked"
    pb.overlap = 100
    pb.gapWidth = 40
    pb.add_data(Reference(ws_a, min_col=4, max_col=5, min_row=p0 + 1, max_row=p_end), titles_from_data=True)
    pb.set_categories(Reference(ws_a, min_col=2, min_row=p0 + 2, max_row=p_end))
    style_chart(pb, "Profit by stock (₹)", 20, max(8, 0.52 * len(by_stock) + 2), legend="b")
    pb.x_axis.scaling.orientation = "maxMin"
    pb.y_axis.crosses = "max"
    pb.y_axis.number_format = "#,##0"
    pb.x_axis.tickLblPos = "low"
    color_series(pb.series[0], GREEN)
    color_series(pb.series[1], RED)
    ws_a.add_chart(pb, f"G{p0}")
    landscape_fit(ws_a)

    wb.active = 0
    wb.save(path)


# ----------------------------------------------------------------------------- optimizer
def write_optimization_report(path, df, meta):
    """df columns: Stop Loss, Take Profit (text like '7%'), Trades, Win Rate (%), Avg Win (%),
    Avg Loss (%), Profit Factor, Sum of Trade Returns (%).  Must already be sorted as desired.
    meta keys: sl_label, tp_label (the pair set in strategy_config.py), train_window, test_window."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Grid_Search"          # FIRST sheet: Dashboard.py reads it by position
    heat = wb.create_sheet("Heatmap")
    ws.sheet_properties.tabColor = BLUE
    heat.sheet_properties.tabColor = NAVY
    cols = list(df.columns)
    n = len(df)
    last = n + 1

    ws.append(cols)
    header_row(ws, 1, 1, len(cols), height=30)
    for i, row in enumerate(df.itertuples(index=False), start=2):
        for c, (name, v) in enumerate(zip(cols, row), start=1):
            cell = ws.cell(i, c, py(v))
            cell.font = F(10)
            cell.border = BOX
            cell.alignment = Alignment(horizontal="center")
            if name == "Profit Factor":
                cell.number_format = "0.00"
            elif name == "Trades":
                cell.number_format = "0"
            elif "%" in name:
                cell.number_format = '0.00"%"'
    for c, name in enumerate(cols, start=1):
        ws.column_dimensions[L(c)].width = max(len(name) + 4, 12)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{L(len(cols))}{last}"
    landscape_fit(ws)
    scale = ColorScaleRule(start_type="min", start_color="F8D0D0", mid_type="percentile", mid_value=50,
                           mid_color="FFF2CC", end_type="max", end_color="CDE7CD")
    for name in ("Win Rate (%)", "Profit Factor", "Sum of Trade Returns (%)"):
        if name in cols:
            letter = L(cols.index(name) + 1)
            ws.conditional_formatting.add(f"{letter}2:{letter}{last}", scale)

    # ---- Heatmap sheet ----
    heat.sheet_view.showGridLines = False
    heat.column_dimensions["A"].width = 14
    for c in range(2, 13):
        heat.column_dimensions[L(c)].width = 11
    heat.column_dimensions["G"].width = 4
    heat.column_dimensions["H"].width = 14
    title_block(heat, 1, "Stop-loss / take-profit grid search",
                f"Tuned on {meta['train_window']}. The chosen pair is then tested on {meta['test_window']}.")

    pf_col, tr_col = L(cols.index("Profit Factor") + 1), L(cols.index("Trades") + 1)
    sl_col, tp_col = L(cols.index("Stop Loss") + 1), L(cols.index("Take Profit") + 1)
    callout = (
        f'="Highest profit factor: SL "&INDEX(Grid_Search!${sl_col}$2:${sl_col}${last},MATCH(MAX(Grid_Search!${pf_col}$2:${pf_col}${last}),Grid_Search!${pf_col}$2:${pf_col}${last},0))'
        f'&" / TP "&INDEX(Grid_Search!${tp_col}$2:${tp_col}${last},MATCH(MAX(Grid_Search!${pf_col}$2:${pf_col}${last}),Grid_Search!${pf_col}$2:${pf_col}${last},0))'
        f'&"  (profit factor "&TEXT(MAX(Grid_Search!${pf_col}$2:${pf_col}${last}),"0.00")&" on "'
        f'&INDEX(Grid_Search!${tr_col}$2:${tr_col}${last},MATCH(MAX(Grid_Search!${pf_col}$2:${pf_col}${last}),Grid_Search!${pf_col}$2:${pf_col}${last},0))&" trades)"'
    )
    c = heat.cell(3, 1, callout)
    c.font = F(11, True, "006300")
    c.fill = fill("E6F4E6")
    heat.merge_cells(start_row=3, start_column=1, end_row=3, end_column=12)
    heat.cell(3, 1).alignment = Alignment(horizontal="left", vertical="center", indent=1)
    heat.row_dimensions[3].height = 24

    key = lambda s: float(str(s).rstrip("%"))
    sls = sorted(df["Stop Loss"].unique(), key=key)
    tps = sorted(df["Take Profit"].unique(), key=key)
    row_of = {(r["Stop Loss"], r["Take Profit"]): i for i, (_, r) in enumerate(df.iterrows(), start=2)}

    def grid(top, left, title, col_name, fmt, kind):
        letter = L(cols.index(col_name) + 1)
        heat.cell(top, left, title).font = F(11, True, NAVY)
        heat.cell(top + 1, left, "SL \\ TP")
        for j, tp in enumerate(tps):
            heat.cell(top + 1, left + 1 + j, tp)
        header_row(heat, top + 1, left, left + len(tps), height=20)
        for i, sl in enumerate(sls):
            r = top + 2 + i
            lab = heat.cell(r, left, sl)
            lab.font = F(10, True)
            lab.fill = fill(CARD)
            lab.alignment = Alignment(horizontal="center", vertical="center")
            lab.border = BOX
            heat.row_dimensions[r].height = 22
            for j, tp in enumerate(tps):
                cell = heat.cell(r, left + 1 + j)
                cell.border = BOX
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = F(10)
                cell.number_format = fmt
                if (sl, tp) in row_of:
                    cell.value = f"=Grid_Search!{letter}{row_of[(sl, tp)]}"
                else:
                    cell.value = "n/a"
                if sl == meta["sl_label"] and tp == meta["tp_label"]:
                    side = Side(style="thick", color=ORANGE)
                    cell.border = Border(left=side, right=side, top=side, bottom=side)
                    cell.font = F(10, True)
        area = f"{L(left + 1)}{top + 2}:{L(left + len(tps))}{top + 1 + len(sls)}"
        vals = pd.to_numeric(df[col_name])
        if kind == "pf" and vals.min() < 1 < vals.max():
            rule = ColorScaleRule(start_type="min", start_color="F2B8B8", mid_type="num", mid_value=1,
                                  mid_color="EEEEEA", end_type="max", end_color="6DA7EC")
        else:
            rule = ColorScaleRule(start_type="min", start_color="E3EEFB", end_type="max", end_color="6DA7EC")
        heat.conditional_formatting.add(area, rule)

    gp, gw = 5, 2 + len(sls) + 3
    grid(gp, 1, "Profit factor", "Profit Factor", "0.00", "pf")
    grid(gp, 8, "Win rate (%)", "Win Rate (%)", '0.0"%"', "seq")
    grid(gp + gw, 1, "Sum of trade returns (%)", "Sum of Trade Returns (%)", '0.0"%"', "seq")
    grid(gp + gw, 8, "Number of trades", "Trades", "0", "seq")

    note_row = gp + 2 * gw
    notes = [
        f"Orange outline = the pair currently set in strategy_config.py (SL {meta['sl_label']} / TP {meta['tp_label']}).",
        "Every cell is a link to the Grid_Search sheet. Pairs with fewer than about 30 trades are statistically weak.",
        "Sum of trade returns adds up per-trade percentages; it is not portfolio growth.",
    ]
    for k, t in enumerate(notes):
        heat.cell(note_row + k, 1, t).font = F(9, False, MUTED, italic=True)
    landscape_fit(heat)

    ws.sheet_view.tabSelected = False   # avoid grouped sheets when Heatmap is the active tab
    wb.active = wb.index(heat)
    heat.sheet_view.tabSelected = True
    wb.save(path)


# ----------------------------------------------------------------------------- screener
def write_screener_report(path, df, meta):
    """df columns: Ticker, LTP (₹), 20-Day SMA (₹), RSI (14), MACD, MACD Signal, Price > SMA,
    RSI Oversold, MACD Bullish, Signal Strength.  meta keys: entry_rsi, exit_rsi, generated."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Market_Screener"
    ws.sheet_properties.tabColor = NAVY
    chart_ws = wb.create_sheet("RSI Chart")
    chart_ws.sheet_properties.tabColor = BLUE
    cols = list(df.columns)
    n = len(df)
    hdr, first = 8, 9
    last = first + n - 1
    ws.sheet_view.showGridLines = False
    title_block(ws, 1, "AlgoEdge Daily Screener",
                f"NIFTY 50  |  Generated {meta['generated']}  |  Entry rule: RSI < {meta['entry_rsi']}")

    for i, name in enumerate(cols, start=1):
        ws.cell(hdr, i, name)
    header_row(ws, hdr, 1, len(cols), height=30)
    sig_fills = {3: "A9D08E", 2: "A9D08E", 1: "FFE699", 0: "F4B084"}
    for r, row in enumerate(df.itertuples(index=False), start=first):
        for c, (name, v) in enumerate(zip(cols, row), start=1):
            v = py(v)
            cell = ws.cell(r, c, v)
            cell.font = F(10, name == "Ticker")
            cell.border = BOX
            if name in ("LTP (₹)", "20-Day SMA (₹)"):
                cell.number_format = RUPEE
            elif name in ("RSI (14)", "MACD", "MACD Signal"):
                cell.number_format = "0.00"
            elif name in ("Price > SMA", "RSI Oversold", "MACD Bullish"):
                cell.alignment = Alignment(horizontal="center")
                ok = bool(v)
                cell.fill = fill("C6EFCE" if ok else "FFC7CE")
                cell.font = F(10, False, "006100" if ok else "9C0006")
            elif name == "Signal Strength":
                cell.alignment = Alignment(horizontal="center")
                cell.font = F(10, True)
                cell.fill = fill(sig_fills.get(int(v), "F4B084"))
    widths = {"Ticker": 14, "LTP (₹)": 16, "20-Day SMA (₹)": 18, "RSI (14)": 14, "MACD": 12,
              "MACD Signal": 14, "Price > SMA": 13, "RSI Oversold": 14, "MACD Bullish": 14, "Signal Strength": 15}
    for i, name in enumerate(cols, start=1):
        ws.column_dimensions[L(i)].width = widths.get(name, 14)
    ws.freeze_panes = f"A{first}"
    ws.auto_filter.ref = f"A{hdr}:{L(len(cols))}{last}"
    rsi_c = L(cols.index("RSI (14)") + 1)
    ws.conditional_formatting.add(
        f"{rsi_c}{first}:{rsi_c}{last}",
        DataBarRule(start_type="num", start_value=0, end_type="num", end_value=100, color=BLUE, showValue=True))

    col = lambda name: f"{L(cols.index(name) + 1)}{first}:{L(cols.index(name) + 1)}{last}"
    cards = [
        ("Stocks scanned", f"=COUNTA({col('Ticker')})", "in the watchlist", "0"),
        (f"Oversold (RSI < {meta['entry_rsi']})", f"=COUNTIF({col('RSI Oversold')},TRUE)",
         "meet the backtest entry rule", "0"),
        ("Strong signals (score >= 2)", f'=COUNTIF({col("Signal Strength")},">=2")', "of 3 conditions met", "0"),
        ("Average RSI", f"=AVERAGE({col('RSI (14)')})", "across the universe", "0.0"),
        ("Lowest RSI", f"=MIN({col('RSI (14)')})",
         f"=INDEX({col('Ticker')},MATCH(MIN({col('RSI (14)')}),{col('RSI (14)')},0))", "0.0"),
    ]
    # 10 columns -> 5 cards of 2 columns, rows 4-6
    for k, (label, val, note, nf) in enumerate(cards):
        card(ws, 4, 1 + 2 * k, 2, label, val, note, number_format=nf)
    ws.row_dimensions[7].height = 10
    landscape_fit(ws)

    # ---- RSI chart sheet ----
    cs = chart_ws
    cs.sheet_view.showGridLines = False
    title_block(cs, 1, "RSI across the universe", "Lower is more oversold. Orange bars would trigger a buy in the backtest.")
    cs["A3"], cs["B3"] = "Entry RSI", meta["entry_rsi"]
    cs["A4"], cs["B4"] = "Exit RSI", meta["exit_rsi"]
    for r in (3, 4):
        cs.cell(r, 1).font = F(10, True)
        cs.cell(r, 2).font = F(10, False, "0000FF")   # blue = input
        cs.cell(r, 2).alignment = Alignment(horizontal="center")
    cs.cell(3, 3, "<- thresholds from strategy_config.py; edit to move the chart lines").font = F(9, False, MUTED, italic=True)
    h2 = 6
    for c, name in enumerate(["Ticker", "RSI (14)", "RSI at/above entry", "Oversold", "Entry line", "Exit line"], start=1):
        cs.cell(h2, c, name)
    header_row(cs, h2, 1, 6, height=30)
    tick_rng = f"Market_Screener!$A${first}:$A${last}"
    rsi_rng = f"Market_Screener!${rsi_c}${first}:${rsi_c}${last}"
    ordered = df.sort_values("RSI (14)")["Ticker"].tolist()
    for k, tk in enumerate(ordered):
        r = h2 + 1 + k
        cs.cell(r, 1, tk).font = F(10, True)
        cs.cell(r, 2, f"=INDEX({rsi_rng},MATCH(A{r},{tick_rng},0))").number_format = "0.0"
        cs.cell(r, 3, f"=IF(B{r}>=$B$3,B{r},0)").number_format = "0.0"
        cs.cell(r, 4, f"=IF(B{r}<$B$3,B{r},0)").number_format = "0.0"
        cs.cell(r, 5, "=$B$3")
        cs.cell(r, 6, "=$B$4")
        for c in range(1, 7):
            cs.cell(r, c).border = BOX
            if c > 1:
                cs.cell(r, c).font = F(10)
                cs.cell(r, c).alignment = Alignment(horizontal="center")
    end = h2 + len(ordered)
    for c, w in zip("ABCDEF", (14, 12, 18, 12, 12, 12)):
        cs.column_dimensions[c].width = w
    bar = BarChart()
    bar.type = "col"
    bar.grouping = "stacked"
    bar.overlap = 100
    bar.gapWidth = 40
    bar.add_data(Reference(cs, min_col=3, max_col=4, min_row=h2, max_row=end), titles_from_data=True)
    bar.set_categories(Reference(cs, min_col=1, min_row=h2 + 1, max_row=end))
    style_chart(bar, "RSI (14) by stock", 32, 11, legend="b")
    bar.y_axis.scaling.min = 0
    bar.y_axis.scaling.max = 100
    color_series(bar.series[0], BLUE)
    color_series(bar.series[1], ORANGE)
    line = LineChart()
    line.add_data(Reference(cs, min_col=5, max_col=6, min_row=h2, max_row=end), titles_from_data=True)
    for s, color in zip(line.series, (MUTED, MUTED)):
        s.graphicalProperties.line.solidFill = color
        s.graphicalProperties.line.width = 15000
        s.smooth = False
    line.series[0].graphicalProperties.line.dashStyle = "dash"
    line.series[1].graphicalProperties.line.dashStyle = "sysDot"
    bar += line
    cs.add_chart(bar, "H3")
    landscape_fit(cs)

    wb.active = 0
    wb.save(path)