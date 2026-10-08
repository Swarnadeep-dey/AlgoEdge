import itertools
import os
import sqlite3
import warnings

import pandas as pd
import yfinance as yf
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from strategy_config import (
    COMMISSION_RATE,
    END_DATE,
    ENTRY_RSI,
    EXIT_RSI,
    RESULTS_DIR,
    RSI_PERIOD,
    START_DATE,
)

# Suppress pandas future warnings for clean output
warnings.simplefilter(action='ignore', category=FutureWarning)

def get_dynamic_tickers(db_name=RESULTS_DIR / "market_data.db"):
    """Pulls unique tickers from the SQLite database with safe fallback logic."""
    fallback_tickers = ["RELIANCE.NS", "TCS.NS", "HDFCBANK.NS", "INFY.NS", "ICICIBANK.NS"]
    
    if not os.path.exists(db_name):
        print(f"Warning: '{db_name}' not found. Using fallback basket.")
        return fallback_tickers
        
    try:
        conn = sqlite3.connect(db_name)
        df = pd.read_sql_query("SELECT DISTINCT Ticker FROM screener_history", conn)
        conn.close()
        
        db_tickers = [str(ticker) + '.NS' for ticker in df['Ticker'].tolist()]
        
        if not db_tickers:
            print("Warning: Database is empty. Using fallback basket.")
            return fallback_tickers
            
        print(f"Successfully loaded {len(db_tickers)} unique tickers from database.")
        return db_tickers
        
    except Exception as e:
        print(f"Database Error: {e}. Using fallback basket.")
        return fallback_tickers

def calculate_rsi(data, window=14):
    delta = data.diff()
    gain = delta.clip(lower=0).ewm(alpha=1/window, adjust=False).mean()
    loss = (-delta.clip(upper=0)).ewm(alpha=1/window, adjust=False).mean()
    rs = gain / loss
    return 100 - (100 / (1 + rs))

def _close_series(raw_data, ticker, single):
    """Extracts one ticker's Close prices from a yfinance frame (works for 1 or many tickers)."""
    if isinstance(raw_data.columns, pd.MultiIndex):
        return raw_data[ticker]['Close']
    return raw_data['Close']

def run_single_backtest(raw_data, tickers, stop_loss_pct, take_profit_pct):
    """Replays the exact Backtester.py rules on pre-downloaded Close prices.

    Entry : RSI < ENTRY_RSI
    Exit  : take-profit, stop-loss (both on the daily close) or RSI > EXIT_RSI
    Return: net of round-trip commission, same formula as Backtester.py
    All thresholds come from strategy_config.py except the SL/TP pair being tested.
    """
    all_trades = []

    for ticker in tickers:
        try:
            closes = _close_series(raw_data, ticker, len(tickers) == 1).ffill().dropna()
            if len(closes) < RSI_PERIOD + 2:
                continue

            rsi = calculate_rsi(closes, RSI_PERIOD)

            in_trade = False
            entry_price = 0.0

            for i in range(RSI_PERIOD, len(closes)):
                price = float(closes.iloc[i])
                latest_rsi = float(rsi.iloc[i])

                if not in_trade:
                    if latest_rsi < ENTRY_RSI:
                        in_trade = True
                        entry_price = price
                else:
                    pnl_pct = (price - entry_price) / entry_price
                    if pnl_pct >= take_profit_pct or pnl_pct <= -stop_loss_pct or latest_rsi > EXIT_RSI:
                        net_return = (price * (1 - COMMISSION_RATE) - entry_price * (1 + COMMISSION_RATE)) / entry_price * 100
                        all_trades.append({'Return (%)': net_return, 'Win': net_return > 0})
                        in_trade = False
        except Exception as e:
            print(f"Skipping {ticker}: {e}")

    return pd.DataFrame(all_trades)

def export_advanced_excel(df, filename=RESULTS_DIR / "Optimization.xlsx"):
    """Exports the optimization results to a professional Excel dashboard with heatmaps."""
    with pd.ExcelWriter(filename, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Grid_Search')
        ws = writer.sheets['Grid_Search']
        
        header_font = Font(bold=True, color='FFFFFF')
        header_fill = PatternFill(start_color='1F4E78', end_color='1F4E78', fill_type='solid')
        thin_border = Border(left=Side(style='thin', color='D4D4D4'), right=Side(style='thin', color='D4D4D4'),
                             top=Side(style='thin', color='D4D4D4'), bottom=Side(style='thin', color='D4D4D4'))

        # Format Headers
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center')
            cell.border = thin_border

        # Auto-adjust column widths and format numbers
        for col_idx, col_name in enumerate(df.columns, 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(max((len(str(v)) for v in df[col_name]), default=0), len(col_name))
            ws.column_dimensions[col_letter].width = max_len + 4
            
            for row_idx in range(2, len(df) + 2):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(horizontal='center')
                
                if '%' in col_name and col_name not in ['Stop Loss', 'Take Profit']:
                    cell.number_format = '0.00"%"'
                elif col_name == 'Profit Factor':
                    cell.number_format = '0.00'

        # Apply Heatmap Color Scales (Green for good, Red for bad)
        last_row = len(df) + 1
        green_red_scale = ColorScaleRule(start_type='min', start_color='FFC7CE',
                                         mid_type='percentile', mid_value=50, mid_color='FFE699',
                                         end_type='max', end_color='C6EFCE')
        
        # Apply heatmaps to Win Rate, Profit Factor, and Total Return
        win_col = get_column_letter(df.columns.get_loc('Win Rate (%)') + 1)
        pf_col = get_column_letter(df.columns.get_loc('Profit Factor') + 1)
        ret_col = get_column_letter(df.columns.get_loc('Total Return (%)') + 1)
        
        ws.conditional_formatting.add(f"{win_col}2:{win_col}{last_row}", green_red_scale)
        ws.conditional_formatting.add(f"{pf_col}2:{pf_col}{last_row}", green_red_scale)
        ws.conditional_formatting.add(f"{ret_col}2:{ret_col}{last_row}", green_red_scale)

    print(f"\nProfessional Heatmap saved to {filename}")

def run_optimizer():
    tickers = get_dynamic_tickers()
    
    stop_losses = [0.03, 0.05, 0.07, 0.10]
    take_profits = [0.05, 0.08, 0.12, 0.15]
    parameter_combinations = list(itertools.product(stop_losses, take_profits))
    
    print(f"Downloading {START_DATE} to {END_DATE} data for {len(tickers)} stocks ONCE...")
    raw_data = yf.download(tickers, start=START_DATE, end=END_DATE, group_by='ticker', threads=True, progress=False)
    
    if raw_data.empty:
        raise SystemExit('[ERROR] No price data downloaded. Check your internet connection and run again.')

    results_summary = []
    print(f"\nRunning {len(parameter_combinations)} backtest iterations...\n")
    
    for sl, tp in parameter_combinations:
        trades_df = run_single_backtest(raw_data, tickers, stop_loss_pct=sl, take_profit_pct=tp)
        
        if not trades_df.empty:
            total_trades = len(trades_df)
            winning_trades = trades_df[trades_df['Win'] == True]
            losing_trades = trades_df[trades_df['Win'] == False]
            
            win_rate = (len(winning_trades) / total_trades) * 100
            total_return = trades_df['Return (%)'].sum()
            
            # New Metrics
            avg_win = winning_trades['Return (%)'].mean() if not winning_trades.empty else 0
            avg_loss = losing_trades['Return (%)'].mean() if not losing_trades.empty else 0
            gross_profit = winning_trades['Return (%)'].sum() if not winning_trades.empty else 0
            gross_loss = abs(losing_trades['Return (%)'].sum()) if not losing_trades.empty else 0
            profit_factor = gross_profit / gross_loss if gross_loss != 0 else float('inf')
            
            results_summary.append({
                'Stop Loss': f"{sl*100:.0f}%",
                'Take Profit': f"{tp*100:.0f}%",
                'Trades': total_trades,
                'Win Rate (%)': round(win_rate, 2),
                'Avg Win (%)': round(avg_win, 2),
                'Avg Loss (%)': round(avg_loss, 2),
                'Profit Factor': round(profit_factor, 2),
                'Total Return (%)': round(total_return, 2)
            })

    results_df = pd.DataFrame(results_summary)
    ranked_results = results_df.sort_values(by='Total Return (%)', ascending=False)
    
    print("================ OPTIMIZATION RESULTS RANKED ================")
    print(ranked_results.to_string(index=False))
    
    # Generate the new Excel Heatmap
    export_advanced_excel(ranked_results)

if __name__ == "__main__":
    run_optimizer()