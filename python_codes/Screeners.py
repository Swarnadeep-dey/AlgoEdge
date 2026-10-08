import socket
import sqlite3
import sys

import pandas as pd
import yfinance as yf
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from strategy_config import ENTRY_RSI, RESULTS_DIR


def internet_ok(host="query1.finance.yahoo.com"):
    """True if Yahoo Finance's hostname resolves (catches DNS / offline / VPN problems early)."""
    try:
        socket.gethostbyname(host)
        return True
    except OSError:
        return False


def run_screener(tickers):
    results = []
    failed = []
    total = len(tickers)
    
    print(f"Starting professional screener for {total} stocks...")
    
    for i, ticker in enumerate(tickers, 1):
        print(f"[{i}/{total}] Analyzing {ticker}...")
        try:
            data = yf.Ticker(ticker).history(period="6mo").dropna(subset=['Close'])
        except Exception as e:
            print(f"   Skipped {ticker}: {type(e).__name__}")
            failed.append(ticker)
            continue
        
        if data.empty or len(data) < 30:
            failed.append(ticker)
            continue
            
        # Technical Calculations
        data['SMA_20'] = data['Close'].rolling(window=20).mean()
        
        delta = data['Close'].diff()
        gain = delta.clip(lower=0).ewm(alpha=1/14, adjust=False).mean()
        loss = (-delta.clip(upper=0)).ewm(alpha=1/14, adjust=False).mean()
        rs = gain / loss
        data['RSI_14'] = 100 - (100 / (1 + rs))
        
        ema_12 = data['Close'].ewm(span=12, adjust=False).mean()
        ema_26 = data['Close'].ewm(span=26, adjust=False).mean()
        data['MACD'] = ema_12 - ema_26
        data['MACD_Signal'] = data['MACD'].ewm(span=9, adjust=False).mean()
        
        # Extract Latest Values
        latest_close = float(data['Close'].iloc[-1])
        latest_sma = float(data['SMA_20'].iloc[-1])
        latest_rsi = float(data['RSI_14'].iloc[-1])
        latest_macd = float(data['MACD'].iloc[-1])
        latest_macd_signal = float(data['MACD_Signal'].iloc[-1])
        
        if any(pd.isna(v) for v in (latest_close, latest_sma, latest_rsi, latest_macd, latest_macd_signal)):
            failed.append(ticker)
            continue
        
        # Boolean Flags
        is_uptrend = latest_close > latest_sma
        is_oversold = latest_rsi < ENTRY_RSI
        is_macd_bullish = latest_macd > latest_macd_signal
        
        # Calculate Signal Strength (0 to 3)
        signal_score = int(is_uptrend) + int(is_oversold) + int(is_macd_bullish)
        clean_ticker = ticker.replace('.NS', '')
        
        results.append({
            'Ticker': clean_ticker,
            'LTP (₹)': latest_close,
            '20-Day SMA (₹)': latest_sma,
            'RSI (14)': latest_rsi,
            'MACD': latest_macd,
            'MACD Signal': latest_macd_signal,
            'Price > SMA': is_uptrend,
            'RSI Oversold': is_oversold,
            'MACD Bullish': is_macd_bullish,
            'Signal Strength': signal_score
        })
        
    if failed:
        print(f"\nWarning: no usable data for {len(failed)} ticker(s): {', '.join(failed)}")
    if not results:
        print("\n[ERROR] No data downloaded for any stock. Existing results were NOT overwritten.")
        print("Check your internet connection / DNS / VPN, then run again.")
        sys.exit(1)

    df_results = pd.DataFrame(results)
    
    # Sort data: Highest signal strength first, then by lowest RSI (most oversold)
    df_results = df_results.sort_values(by=['Signal Strength', 'RSI (14)'], ascending=[False, True])
    # SQLite database storage
    conn = sqlite3.connect(RESULTS_DIR / "market_data.db")
    df_results.to_sql("screener_history", conn, if_exists="replace", index=False)
    conn.close()

    output_file = RESULTS_DIR / "Screener_Results.xlsx"
    
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        df_results.to_excel(writer, index=False, sheet_name='Market_Screener')
        
        workbook = writer.book
        worksheet = writer.sheets['Market_Screener']
        
        # --- Corporate Styling Definitions ---
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        
        true_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
        true_font = Font(color="006100")
        false_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
        false_font = Font(color="9C0006")
        
        # Signal Strength Colors
        strong_fill = PatternFill(start_color="A9D08E", end_color="A9D08E", fill_type="solid")
        medium_fill = PatternFill(start_color="FFE699", end_color="FFE699", fill_type="solid")
        weak_fill = PatternFill(start_color="F4B084", end_color="F4B084", fill_type="solid")
        
        thin_border = Border(left=Side(style='thin', color='D4D4D4'),
                             right=Side(style='thin', color='D4D4D4'),
                             top=Side(style='thin', color='D4D4D4'),
                             bottom=Side(style='thin', color='D4D4D4'))
        
        rupee_format = '[$₹-en-IN]#,##0.00'
        decimal_format = '0.00'
        
        # Format Headers
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = thin_border
            
        # Freeze top row and add filters
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        
        # Format Data Rows
        for row_idx in range(2, len(df_results) + 2):
            for col_idx, col_name in enumerate(df_results.columns, 1):
                cell = worksheet.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                
                # Number Formatting
                if col_name in ['LTP (₹)', '20-Day SMA (₹)']:
                    cell.number_format = rupee_format
                elif col_name in ['RSI (14)', 'MACD', 'MACD Signal']:
                    cell.number_format = decimal_format
                    
                # Boolean Conditional Formatting
                if col_name in ['Price > SMA', 'RSI Oversold', 'MACD Bullish']:
                    cell.alignment = Alignment(horizontal='center')
                    if cell.value == True:
                        cell.fill = true_fill
                        cell.font = true_font
                    elif cell.value == False:
                        cell.fill = false_fill
                        cell.font = false_font
                
                # Signal Strength Heatmap (FIXED)
                if col_name == 'Signal Strength':
                    cell.alignment = Alignment(horizontal='center')
                    cell.font = Font(bold=True)  # Applied bold using the Font class
                    if cell.value >= 2:
                        cell.fill = strong_fill
                    elif cell.value == 1:
                        cell.fill = medium_fill
                    else:
                        cell.fill = weak_fill

        # Dynamic Column Widths
        for col_idx, col_name in enumerate(df_results.columns, 1):
            col_letter = get_column_letter(col_idx)
            max_length = max(max((len(str(v)) for v in df_results[col_name]), default=0), len(col_name))
            worksheet.column_dimensions[col_letter].width = max_length + 4
                        
    print(f"\nAnalysis complete. Executive dashboard saved to {output_file}")

# Watchlist
nifty_50_tickers = [
    'ADANIENT.NS', 'ADANIPORTS.NS', 'APOLLOHOSP.NS', 'ASIANPAINT.NS', 'AXISBANK.NS',
    'BAJAJ-AUTO.NS', 'BAJFINANCE.NS', 'BAJAJFINSV.NS', 'BPCL.NS', 'BHARTIARTL.NS',
    'BRITANNIA.NS', 'CIPLA.NS', 'COALINDIA.NS', 'DIVISLAB.NS', 'DRREDDY.NS',
    'EICHERMOT.NS', 'GRASIM.NS', 'HCLTECH.NS', 'HDFCBANK.NS', 'HDFCLIFE.NS',
    'HEROMOTOCO.NS', 'HINDALCO.NS', 'HINDUNILVR.NS', 'ICICIBANK.NS', 'ITC.NS',
    'INDUSINDBK.NS', 'INFY.NS', 'JSWSTEEL.NS', 'KOTAKBANK.NS', 'LT.NS',
    'TRENT.NS', 'M&M.NS', 'MARUTI.NS', 'NESTLEIND.NS', 'NTPC.NS',
    'ONGC.NS', 'POWERGRID.NS', 'RELIANCE.NS', 'SBILIFE.NS', 'SBIN.NS',
    'SUNPHARMA.NS', 'TATACONSUM.NS', 'TMPV.NS', 'TATASTEEL.NS', 'TCS.NS',
    'TECHM.NS', 'TITAN.NS', 'ULTRACEMCO.NS', 'WIPRO.NS', 'SHRIRAMFIN.NS'
]

if __name__ == "__main__":
    if not internet_ok():
        print("[ERROR] Cannot reach Yahoo Finance (DNS lookup failed).")
        print("Check your internet connection, VPN/proxy or firewall, then run again.")
        sys.exit(1)
    run_screener(nifty_50_tickers)